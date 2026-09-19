from copy import deepcopy
from pathlib import Path
from typing import Any

from pptx import Presentation


def replace_shape_text(
    shape: Any,
    lines: list[str],
) -> None:
    text_frame = shape.text_frame
    first_paragraph = text_frame.paragraphs[0]

    paragraph_template = deepcopy(first_paragraph._p)
    run_properties = (
        deepcopy(first_paragraph.runs[0]._r.rPr) if first_paragraph.runs else None
    )

    text_frame.clear()

    for index, line in enumerate(lines):
        if index == 0:
            paragraph = text_frame.paragraphs[0]
        else:
            text_frame._txBody.append(
                deepcopy(
                    paragraph_template,
                )
            )
            paragraph = text_frame.paragraphs[-1]

        paragraph.clear()
        run = paragraph.add_run()

        if run_properties is not None:
            run._r.insert(
                0,
                deepcopy(run_properties),
            )

        run.text = line


def fill_seating_blocks(
    template_path: Path | str,
    output_path: Path | str,
    selected_blocks: list[dict[str, Any]],
    tables: list[list[str]],
) -> None:

    if len(selected_blocks) < len(tables):
        raise ValueError(
            "Недостаточно блоков для вставки рассадки",
        )

    prs = Presentation(
        str(template_path),
    )

    for table_number, (block, names) in enumerate(
        zip(
            selected_blocks,
            tables,
        ),
        start=1,
    ):
        slide = prs.slides[block["slide"] - 1]

        shape = next(
            (shape for shape in slide.shapes if shape.shape_id == block["shape_id"]),
            None,
        )

        if shape is None:
            raise ValueError(
                f"Shape не найден: "
                f"slide={block['slide']}, "
                f"shape_id={block['shape_id']}"
            )

        if not shape.has_text_frame:
            raise ValueError(
                f"Shape не содержит текст: " f"shape_id={block['shape_id']}"
            )

        replace_shape_text(
            shape,
            [
                f"Стол {table_number}",
                *names,
            ],
        )

    output = Path(output_path) if isinstance(output_path, str) else output_path
    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    prs.save(
        str(output),
    )
