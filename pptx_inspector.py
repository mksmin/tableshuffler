import logging
import time
from collections import defaultdict
from typing import Any

from pptx import Presentation

from ai_analyzer import find_seating_groups
from pptx_writer import fill_seating_blocks
from validation import validate_seating_block

log = logging.getLogger(__name__)


def group_blocks_by_text(
    text_blocks: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)

    for block in text_blocks:
        key = block["text"].strip().lower()
        groups[key].append(block)

    result = []

    for group_id, (text, blocks) in enumerate(
        groups.items(),
        start=1,
    ):

        result.append(
            {
                "group_id": group_id,
                "text": text,
                "count": len(blocks),
                "shape_ids": [block["shape_id"] for block in blocks],
                "blocks": blocks,
            },
        )
    return result


def select_seating_blocks(
    text_blocks: list[dict[str, Any]],
    seating_blocks: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Возвращает найденные блоки в исходном порядке объектов PowerPoint."""
    seating_keys = {(block["slide"], block["shape_id"]) for block in seating_blocks}
    return [
        block
        for block in text_blocks
        if (block["slide"], block["shape_id"]) in seating_keys
    ]


def inspect_pptx(
    pptx_path: str,
) -> list[dict[str, Any]]:
    prs = Presentation(pptx_path)

    pres_len = len(prs.slides)

    slide_width = prs.slide_width
    slide_height = prs.slide_height

    text_blocks = []
    choice_slide = 0

    if pres_len > 1:
        while True:
            try:
                choice_slide = int(
                    input(
                        "Введи номер слайда: ",
                    )
                )
            except ValueError:
                log.info("Введи номер слайда числом")
                continue

            if choice_slide <= 0 or choice_slide > pres_len:
                log.info(
                    "Ошибка, выбери номер слайда. Всего в презентации слайдов: %s",
                    pres_len,
                )
                continue
            else:
                choice_slide -= 1
                break

    slide = prs.slides[choice_slide]

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue

        text_blocks.append(
            {
                "slide": choice_slide + 1,
                "shape_id": shape.shape_id,
                "name": shape.name,
                "text": shape.text.strip(),
                "x": round(shape.left / slide_width * 100, 1),
                "y": round(shape.top / slide_height * 100, 1),
                "width": round(shape.width / slide_width * 100, 1),
                "height": round(shape.height / slide_height * 100, 1),
            }
        )

    return text_blocks


if __name__ == "__main__":
    logging.basicConfig(
        level=logging.INFO,
        format="%(message)s",
    )
    start = time.time()
    pptx_path = "./footages/template.pptx"

    text_bl = inspect_pptx(pptx_path)
    groups_block = group_blocks_by_text(text_bl)

    llm_start = time.time()
    seating_groups = find_seating_groups(groups_block)
    llm_end = time.time()
    llm_duration = llm_end - llm_start

    required_tables = 4
    seating_blocks = validate_seating_block(
        seating_groups,
        required_tables,
    )

    log.info(
        "Валидация пройдена: требуется %s, доступно %s",
        required_tables,
        len(seating_blocks),
    )

    selected_blocks = select_seating_blocks(
        text_bl,
        seating_blocks,
    )[:required_tables]

    log.debug("Порядок блоков:")

    for table_num, block in enumerate(
        selected_blocks,
        start=1,
    ):
        log.debug(
            "Стол %s: slide=%s, shape_id=%s, x=%s, y=%s",
            table_num,
            block["slide"],
            block["shape_id"],
            block["x"],
            block["y"],
        )

    end = time.time()
    duration = end - start

    log.debug("duration: %.3f s", duration)
    log.debug("llm_duration: %.3f s", llm_duration)

    tables = [
        [
            "Иванов Иван Иванович",
            "Петров Петр Петрович",
        ],
        [
            "Иванов Иван Иванович",
            "Петров Петр Петрович",
        ],
        [
            "Иванов Иван Иванович",
            "Петров Петр Петрович",
        ],
        [
            "Иванов Иван Иванович",
            "Петров Петр Петрович",
        ],
    ]

    fill_seating_blocks(
        template_path=pptx_path,
        output_path="./footages/result.pptx",
        selected_blocks=selected_blocks,
        tables=tables,
    )
